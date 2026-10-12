#!/usr/bin/env python3
"""
Ingest the ReadLark launch plan (readlark-top-100-books.xlsx) into a
machine-readable queue the daily factory reads from.

Input : readlark-top-100-books.xlsx  (sheets "First 20 Launch Plan" + "Top 100 Books")
Output: content/launch-plan.json  — ordered list of:
        { order, slug, title, author, year, category, week, status, produced_date }

status = "done" if a script already exists at content/videos/<slug>.json, else "queued".
category is NORMALIZED onto the curated CATEGORIES set (schema.ts) via keywords.

Usage: python3 pipeline/ingest_launch_plan.py
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
XLSX = REPO / "readlark-top-100-books.xlsx"


def slugify(title: str) -> str:
    s = title.lower().replace("&", "and")
    s = re.sub(r"[''`]", "", s)          # drop apostrophes (can't -> cant)
    s = re.sub(r"[^a-z0-9]+", "-", s)     # everything else -> hyphen
    return re.sub(r"-+", "-", s).strip("-")


def normalize_category(cat: str) -> str:
    """Map the xlsx's loose categories onto the curated playlist set."""
    c = (cat or "").lower()
    rules = [
        (("money", "finance", "wealth", "invest"), "Money & Investing"),
        (("habit", "productivity"), "Habits & Productivity"),
        (("negoti", "communicat"), "Communication"),
        (("philosoph",), "Philosophy"),
        (("psycholog", "power"), "Psychology"),
        (("motivat", "disciplin"), "Motivation & Discipline"),
        (("health",), "Health & Fitness"),
        (("purpose", "ikigai", "meaning"), "Purpose & Meaning"),
    ]
    for keys, label in rules:
        if any(k in c for k in keys):
            return label
    return "Mindset & Self-Help"  # self-help / spirituality / growth / default


def parse_int(v):
    if v is None:
        return None
    m = re.search(r"\d{3,4}", str(v))
    return int(m.group()) if m else None


def main() -> None:
    if not XLSX.exists():
        sys.exit(f"ERROR: launch plan not found: {XLSX}")
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # Metadata lookup from "Top 100 Books": title -> {author, year, category}
    meta = {}
    for r in wb["Top 100 Books"].iter_rows(min_row=5, values_only=True):
        if not r or not r[1]:
            continue
        meta[str(r[1]).strip().lower()] = {"author": r[2], "year": r[3], "category": r[4]}

    plan = []
    for r in wb["First 20 Launch Plan"].iter_rows(min_row=5, values_only=True):
        if not r or r[0] is None or not r[1]:
            continue
        title = str(r[1]).strip()
        m = meta.get(title.lower(), {})
        slug = slugify(title)
        status = "done" if (REPO / "content" / "videos" / f"{slug}.json").exists() else "queued"
        plan.append({
            "order": int(float(r[0])),
            "slug": slug,
            "title": title,
            "author": (str(m.get("author")).strip() if m.get("author") else None),
            "year": parse_int(m.get("year")),
            "category": normalize_category(m.get("category", "")),
            "week": str(r[2]).strip() if r[2] else None,
            "status": status,
            "produced_date": None,
        })

    plan.sort(key=lambda b: b["order"])
    out = REPO / "content" / "launch-plan.json"
    out.write_text(json.dumps(plan, indent=2) + "\n")
    done = sum(1 for b in plan if b["status"] == "done")
    print(f"✓ launch-plan.json: {len(plan)} books ({done} done, {len(plan)-done} queued) → {out}")
    for b in plan:
        print(f"  {b['order']:>2} [{b['status']:>6}] {b['slug']:36} {b['category']}")


if __name__ == "__main__":
    main()
