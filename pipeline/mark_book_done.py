#!/usr/bin/env python3
"""
Mark a book as produced: writes a "Video" column (= produced date) into BOTH
sheets of readlark-top-100-books.xlsx, and flips the book's status→done in
content/launch-plan.json. Idempotent (re-running updates in place, no dup column).

Usage: python3 pipeline/mark_book_done.py <slug> [YYYY-MM-DD]
"""
import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
XLSX = REPO / "readlark-top-100-books.xlsx"
PLAN = REPO / "content" / "launch-plan.json"
VIDEO_HEADER = "Video"


def slugify(title: str) -> str:
    s = title.lower().replace("&", "and")
    s = re.sub(r"[''`]", "", s)
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return re.sub(r"-+", "-", s).strip("-")


def mark_sheet(ws, slug: str, when: str) -> bool:
    """Find the header row + title column, ensure a 'Video' column, set the row."""
    header_row = None
    title_col = None
    for r in range(1, min(ws.max_row, 10) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip() in ("Title", "Book"):
                header_row, title_col = r, c
                break
        if header_row:
            break
    if not header_row:
        return False
    # find or create the Video column
    video_col = None
    for c in range(1, ws.max_column + 2):
        if (ws.cell(header_row, c).value or "") == VIDEO_HEADER:
            video_col = c
            break
    if not video_col:
        video_col = ws.max_column + 1
        ws.cell(header_row, video_col, VIDEO_HEADER)
    # find the book row by slugified title
    for r in range(header_row + 1, ws.max_row + 1):
        t = ws.cell(r, title_col).value
        if t and slugify(str(t)) == slug:
            ws.cell(r, video_col, when)
            return True
    return False


def main() -> None:
    if len(sys.argv) < 2:
        sys.exit("Usage: python3 pipeline/mark_book_done.py <slug> [YYYY-MM-DD]")
    slug = sys.argv[1]
    when = sys.argv[2] if len(sys.argv) > 2 else date.today().isoformat()

    # launch-plan.json
    if PLAN.exists():
        plan = json.loads(PLAN.read_text())
        hit = next((b for b in plan if b["slug"] == slug), None)
        if hit:
            hit["status"] = "done"
            hit["produced_date"] = when
            PLAN.write_text(json.dumps(plan, indent=2) + "\n")
            print(f"✓ launch-plan.json: {slug} → done ({when})")
        else:
            print(f"⚠ {slug} not in launch-plan.json")

    # xlsx both sheets
    if XLSX.exists():
        wb = openpyxl.load_workbook(XLSX)
        marked = [ws.title for ws in wb.worksheets if mark_sheet(ws, slug, when)]
        wb.save(XLSX)
        print(f"✓ xlsx: marked '{VIDEO_HEADER}={when}' on sheets: {', '.join(marked) or '(no matching row)'}")
    else:
        print(f"⚠ xlsx not found: {XLSX}")


if __name__ == "__main__":
    main()
